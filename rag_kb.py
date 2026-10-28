"""Lesson 26: local document indexing and BM25-style retrieval.

Supported: .txt, .md, .pdf (text layer), .docx. No network or model calls.
Database is separate from the existing multiagent.db.
"""
from __future__ import annotations

import argparse
from collections import Counter
import hashlib
import math
from pathlib import Path
import re
import sqlite3
import sys

ROOT = Path(__file__).resolve().parent
DEFAULT_DOCS = ROOT / "knowledge_docs"
DEFAULT_DB = ROOT / "knowledge_base.db"
SUPPORTED = {".txt", ".md", ".pdf", ".docx"}


def tokenize(text: str) -> list[str]:
    """Language-neutral tokens; Chinese character n-grams + Latin terms."""
    text = text.lower()
    tokens = []
    for part in re.findall(r"[a-z0-9]+|[\u3400-\u9fff]+", text):
        if '\u3400' <= part[0] <= '\u9fff':
            tokens.extend(part[i:i + 2] for i in range(len(part) - 1))
            if len(part) == 1:
                tokens.append(part)
        else:
            tokens.append(part)
    return tokens


def init_db(db_path: Path = DEFAULT_DB) -> None:
    db_path.parent.mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(db_path, timeout=30) as db:
        db.execute("PRAGMA busy_timeout=30000")
        db.execute("""
            CREATE TABLE IF NOT EXISTS kb_documents (
                id INTEGER PRIMARY KEY,
                source_path TEXT NOT NULL UNIQUE,
                sha256 TEXT NOT NULL,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            )
        """)
        db.execute("""
            CREATE TABLE IF NOT EXISTS kb_chunks (
                id INTEGER PRIMARY KEY,
                document_id INTEGER NOT NULL REFERENCES kb_documents(id),
                source_label TEXT NOT NULL,
                chunk_index INTEGER NOT NULL,
                content TEXT NOT NULL,
                UNIQUE(document_id, chunk_index)
            )
        """)
        db.execute("CREATE INDEX IF NOT EXISTS kb_chunks_document_idx ON kb_chunks(document_id)")


def extract_sections(path: Path) -> list[tuple[str, str]]:
    """Return (location label, text). Preserve PDF page labels."""
    ext = path.suffix.lower()
    if ext in {'.txt', '.md'}:
        content = path.read_text(encoding='utf-8-sig')
        return [('全文', content)]
    if ext == '.pdf':
        try:
            import fitz
        except ImportError as exc:
            raise RuntimeError('缺少 PyMuPDF：python -m pip install pymupdf') from exc
        with fitz.open(path) as doc:
            return [(f'第{i + 1}页', page.get_text('text')) for i, page in enumerate(doc)]
    if ext == '.docx':
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError('缺少 python-docx：python -m pip install python-docx') from exc
        doc = Document(path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    paragraphs.append(' | '.join(cells))
        return [('正文及表格', '\n'.join(paragraphs))]
    raise ValueError(f'不支持的文件格式: {path.suffix}')


def split_text(text: str, size: int = 500, overlap: int = 80) -> list[str]:
    """Character-based chunks; overlap retains context at boundaries."""
    text = re.sub(r'\s+', ' ', text).strip()
    if not text:
        return []
    if size <= 0 or not 0 <= overlap < size:
        raise ValueError('必须满足 size > 0 且 0 <= overlap < size')
    result = []
    step = size - overlap
    for start in range(0, len(text), step):
        piece = text[start:start + size]
        if piece:
            result.append(piece)
        if start + size >= len(text):
            break
    return result


def index_directory(folder: Path = DEFAULT_DOCS, db_path: Path = DEFAULT_DB) -> dict:
    folder = folder.resolve()
    if not folder.is_dir():
        raise FileNotFoundError(f'文档目录不存在：{folder}')
    init_db(db_path)
    counts = {'added': 0, 'updated': 0, 'unchanged': 0, 'skipped': 0, 'chunks': 0}
    paths = sorted(p for p in folder.rglob('*') if p.is_file() and p.suffix.lower() in SUPPORTED and not p.name.startswith('~$'))
    with sqlite3.connect(db_path, timeout=30) as db:
        db.execute('PRAGMA foreign_keys=ON')
        for path in paths:
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            relative = path.relative_to(folder).as_posix()
            old = db.execute('SELECT id, sha256 FROM kb_documents WHERE source_path=?', (relative,)).fetchone()
            if old and old[1] == digest:
                counts['unchanged'] += 1
                continue
            try:
                sections = extract_sections(path)
                chunks = [(location, chunk) for location, text in sections for chunk in split_text(text)]
            except Exception as exc:
                counts['skipped'] += 1
                print(f'[跳过] {relative}: {exc}', file=sys.stderr)
                continue
            if not chunks:
                counts['skipped'] += 1
                print(f'[跳过] {relative}: 未提取到文字（扫描版 PDF 需要 OCR）', file=sys.stderr)
                continue
            if old:
                doc_id = old[0]
                db.execute('DELETE FROM kb_chunks WHERE document_id=?', (doc_id,))
                db.execute('UPDATE kb_documents SET sha256=?, updated_at=CURRENT_TIMESTAMP WHERE id=?', (digest, doc_id))
                counts['updated'] += 1
            else:
                cur = db.execute('INSERT INTO kb_documents(source_path,sha256) VALUES (?,?)', (relative, digest))
                doc_id = cur.lastrowid
                counts['added'] += 1
            db.executemany('''INSERT INTO kb_chunks(document_id,source_label,chunk_index,content)
                              VALUES (?,?,?,?)''', [(doc_id, loc, i + 1, content) for i, (loc, content) in enumerate(chunks)])
            counts['chunks'] += len(chunks)
            print(f'[已索引] {relative}: {len(chunks)} 个片段')
    return counts


def search(query: str, *, top_k: int = 5, db_path: Path = DEFAULT_DB) -> list[dict]:
    if not db_path.exists():
        raise FileNotFoundError('知识库不存在，请先执行 index 命令')
    query_terms = set(tokenize(query))
    if not query_terms:
        return []
    with sqlite3.connect(db_path) as db:
        rows = db.execute('''SELECT c.id,d.source_path,c.source_label,c.chunk_index,c.content
                             FROM kb_chunks c JOIN kb_documents d ON d.id=c.document_id''').fetchall()
    if not rows:
        return []
    counts = [Counter(tokenize(r[4])) for r in rows]
    doc_freq = Counter(term for c in counts for term in c)
    avg_len = max(sum(sum(c.values()) for c in counts) / len(counts), 1)
    n = len(rows)
    results = []
    for row, tf in zip(rows, counts):
        length = sum(tf.values())
        score = 0.0
        for term in query_terms:
            frequency = tf.get(term, 0)
            if not frequency:
                continue
            df = doc_freq[term]
            idf = math.log(1 + (n - df + 0.5) / (df + 0.5))
            score += idf * frequency * 2.2 / (frequency + 1.2 * (0.25 + 0.75 * length / avg_len))
        if score > 0:
            results.append({'chunk_id': row[0], 'source': row[1], 'location': row[2],
                            'chunk_index': row[3], 'content': row[4], 'score': score})
    results.sort(key=lambda item: (-item['score'], item['source'], item['chunk_index']))
    return results[:max(top_k, 0)]


def main() -> None:
    parser = argparse.ArgumentParser(description='第26课：本地企业知识库')
    parser.add_argument('--db', type=Path, default=DEFAULT_DB, help='SQLite 知识库文件')
    sub = parser.add_subparsers(dest='command', required=True)
    p_index = sub.add_parser('index', help='导入文档并创建索引')
    p_index.add_argument('--folder', type=Path, default=DEFAULT_DOCS)
    p_search = sub.add_parser('search', help='检索相关文档片段')
    p_search.add_argument('query', help='检索问题')
    p_search.add_argument('--top-k', type=int, default=5)
    args = parser.parse_args()
    if args.command == 'index':
        counts = index_directory(args.folder, args.db)
        print('索引完成：', counts)
    elif args.command == 'search':
        hits = search(args.query, top_k=args.top_k, db_path=args.db)
        if not hits:
            print('未找到相关内容。请检查文档是否已导入，或更换关键词。')
        for i, hit in enumerate(hits, 1):
            print(f'\n[{i}] {hit["source"]} · {hit["location"]} · 片段{hit["chunk_index"]} · 分数{hit["score"]:.3f}')
            print(hit['content'][:600])


if __name__ == '__main__':
    main()
